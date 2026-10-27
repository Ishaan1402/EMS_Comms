"""Score the AI against what happened to the patient, not only against a triage label (KTAS set).

    python -m evals.benchmark.outcomes --data evals/local/benchmark_data   # reads the KTAS source once
    python -m evals.benchmark.outcomes                                     # rescore from the saved file

The KTAS source records each visit's disposition. A triage level is a judgement, even the experts'
(they rated only about half of the patients who went to ICU, died or had surgery from the ED as 1-2),
so this asks the outcome question directly: of the patients with a critical outcome, how many did
each rater put at 1-2, and of those sent home, how many did each one flag as urgent?

Writes results/ktas_dispositions.csv (case and disposition, from the source's S1 Appendix, CC BY
4.0) and results/outcomes.json (every count with its 95% Wilson interval, primary run). to_sas.py
loads the same counts into CAS as EIR_BENCH_OUTCOMES. No model calls.
"""
import argparse
import csv
import json
from pathlib import Path

from .stats import wilson

HERE = Path(__file__).parent
RESULTS = HERE / "results"
DISPOSITIONS = RESULTS / "ktas_dispositions.csv"
SOURCE = Path("ktas/S1_EMR_review_coding_data.xlsx")
# The source's coding guideline (S2 Appendix).
DISPOSITION_LABELS = {"1": "Discharge", "2": "Ward admission", "3": "ICU admission", "4": "AMA discharge",
                      "5": "Transfer", "6": "Death", "7": "Surgery from ED"}
GROUPS = {
    "ICU, death or surgery from ED": {"3", "6", "7"},
    "Ward admission": {"2"},
    "Sent home (discharge or AMA)": {"1", "4"},
}
SCOPES = {"KTAS all visits": lambda r: True, "KTAS ambulance arrivals": lambda r: bool(r["EMS_ARRIVAL"])}
NO_LEVEL = "No level (Needs review)"
RATERS = {
    "Eir AI": lambda r: NO_LEVEL if not r["AI_GAVE_LEVEL"] else "Given 1-2" if r["AI_URGENT"] else "Given 3-5",
    "Triage nurses": lambda r: "Given 1-2" if r["NURSE_KTAS"] <= 2 else "Given 3-5",
    "KTAS experts": lambda r: "Given 1-2" if r["REFERENCE_LEVEL"] <= 2 else "Given 3-5",
    "Simple vitals rule": lambda r: "Given 1-2" if r["SIMPLE_RULE"] <= 2 else "Given 3-5",
}


def read_source(data: Path) -> dict:
    """{case_id: disposition code} from the KTAS spreadsheet, matched by its sheet row as build_cases does."""
    from openpyxl import load_workbook

    from .run_live import load_cases
    sheet = load_workbook(data / SOURCE, read_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    column = rows[0].index("Disposition")
    by_row = {number: str(row[column]) for number, row in enumerate(rows[1:], start=2)}
    return {case["case_id"]: by_row[case["meta"]["source_row"]["sheet_row"]] for case in load_cases("ktas")}


def read_dispositions(path: Path = DISPOSITIONS) -> dict:
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as file:
        return {row["case_id"]: row["disposition"] for row in csv.DictReader(file)}


def outcome_rows(calls: list, dispositions: dict) -> list:
    """Long rows (scope, outcome group, rater, placed at, visits) for the primary KTAS run."""
    ktas = [r for r in calls if r["DATASET"] == "ktas" and r["PRIMARY"] and r["CASE_ID"] in dispositions]
    rows = []
    for scope, chosen in SCOPES.items():
        for group, codes in GROUPS.items():
            visits = [r for r in ktas if chosen(r) and dispositions[r["CASE_ID"]] in codes]
            for rater, placed_by in RATERS.items():
                placed = [placed_by(r) for r in visits]
                for outcome in ("Given 1-2", NO_LEVEL, "Given 3-5"):
                    if outcome == NO_LEVEL and rater != "Eir AI":
                        continue
                    rows.append({"SCOPE": scope, "OUTCOME_GROUP": group, "RATER": rater, "PLACED": outcome,
                                 "VISITS": placed.count(outcome), "GROUP_VISITS": len(visits),
                                 "SHARE": round(placed.count(outcome) / len(visits), 4) if visits else None})
    return rows


def summary(rows: list) -> dict:
    """{scope: {group: {rater: {placed: wilson}}}} for outcomes.json."""
    out = {}
    for row in rows:
        out.setdefault(row["SCOPE"], {}).setdefault(row["OUTCOME_GROUP"], {}).setdefault(row["RATER"], {})[
            row["PLACED"]] = wilson(row["VISITS"], row["GROUP_VISITS"])
    return out


def main(argv=None) -> int:
    from .to_sas import build_tables
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--data", type=Path, help="folder holding ktas/S1_EMR_review_coding_data.xlsx")
    args = parser.parse_args(argv)
    if args.data:
        dispositions = read_source(args.data)
        with DISPOSITIONS.open("w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file, lineterminator="\n")
            writer.writerow(["case_id", "disposition", "disposition_label"])
            for case_id, code in dispositions.items():
                writer.writerow([case_id, code, DISPOSITION_LABELS.get(code, "")])
        print(f"{DISPOSITIONS}: {len(dispositions)} visits")
    rows = build_tables()["OUTCOMES"]
    (RESULTS / "outcomes.json").write_text(json.dumps(summary(rows), indent=1) + "\n", encoding="utf-8")
    for row in rows:
        if row["SCOPE"] == "KTAS all visits":
            print(f"  {row['OUTCOME_GROUP']:30} {row['RATER']:18} {row['PLACED']:24} {row['VISITS']}/{row['GROUP_VISITS']}")
    print(f"{RESULTS / 'outcomes.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
