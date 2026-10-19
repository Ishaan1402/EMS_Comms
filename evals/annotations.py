"""Read the existing vertical clinician form without modifying the workbook."""

import hashlib
import re
from pathlib import Path

from .contracts import CHANGE, PREPARATION, STATUSES, blank_reference, case_dataset_version, text

PATIENT_ANSWERS = {
    "EMS to ED Handoff": "reference_summary",
    "Meaningful change?": "meaningful_change",
    "Direction (if Yes)": "change_direction",
    "Evidence / missing information": "change_evidence",
    "Preparation needed?": "preparation_category",
    "Reason for preparation / limits": "preparation_explanation",
    "Review status": "status",
}
PATIENT_PREPARATION = ("Prepare now", "Can wait", "Routine", "Cannot assess")
PATIENT_DIRECTIONS = ("Improvement", "Deterioration", "Mixed", "Unclear")


def parse_patient_rows(rows, cases):
    """Import the simple patient-review form against the exact versioned report fixtures.

    Earlier/latest stays a two-report question for the clinician. The existing replay supplies
    its acknowledgment setup; no additional acknowledgment fields are collected here.
    """
    version = case_dataset_version(cases)
    if not version:
        raise ValueError("Patient review needs versioned case fixtures; select the matching --cases file")
    known = {case["case_id"]: case for case in cases}
    references, active, seen, metadata = [], None, set(), {}

    def finish():
        if active is None:
            return
        case = known[active["case_id"]]
        expected = set(PATIENT_ANSWERS.values()) | {"patient_details", "current_transcript"}
        if case["prior_information"] != "No earlier report provided.":
            expected.add("prior_information")
        if seen != expected:
            raise ValueError(f"Incomplete/changed Patient review layout for {active['case_id']}")
        active["explanation_missing_information"] = "\n".join(
            f"{label}: {active[field]}" for label, field in (
                ("Change / missing information", "change_evidence"),
                ("Preparation / limits", "preparation_explanation")) if active.get(field)) or None
        for field, choices in (("meaningful_change", CHANGE), ("change_direction", PATIENT_DIRECTIONS),
                               ("status", STATUSES)):
            if active[field] is not None and active[field] not in choices:
                raise ValueError(f"Invalid {field} for {active['case_id']}: {active[field]}")
        display = active["preparation_category"]
        if display is not None and display not in PATIENT_PREPARATION:
            raise ValueError(f"Invalid preparation_category for {active['case_id']}: {display}")
        active["preparation_category_display"] = display
        if display == "Cannot assess":
            active["preparation_category"] = "Unsure"

    for row in rows:
        label, value = text(row[0] if row else None), text(row[1] if len(row) > 1 else None)
        match = re.fullmatch(r"Case \d+ / (SYN\d+)", label)
        if match:
            finish()
            case_id = match.group(1)
            if case_id not in known or any(ref["case_id"] == case_id for ref in references):
                raise ValueError(f"Unknown/duplicate case card: {case_id}")
            if value != f"ETA {known[case_id]['eta_minutes']} minutes at latest report":
                raise ValueError(f"ETA changed for {case_id}; version cases first")
            active, seen = blank_reference(case_id), set()
            references.append(active)
            continue
        if active is None:
            if label == "Reviewer alias":
                metadata[label] = value
            if label == "Dataset version":
                if label in metadata:
                    raise ValueError("Duplicate Dataset version")
                metadata[label] = value
            continue
        if not label and not value:
            continue
        case = known[active["case_id"]]
        if label in PATIENT_ANSWERS:
            field = PATIENT_ANSWERS[label]
        elif label == "Age / reported sex / weight":
            field = "patient_details"
        elif re.fullmatch(r"T=0\n(?:Updated )?EMS Field Report", label):
            field = "current_transcript"
        elif re.fullmatch(r"T-\d+ min\nEMS Field Report", label):
            if label != f"T-{case['elapsed_minutes']} min\nEMS Field Report":
                raise ValueError(f"Report interval changed for {active['case_id']}; version cases first")
            field = "prior_information"
        else:
            raise ValueError(f"Unexpected Patient review field for {active['case_id']}: {label}")
        if field in seen:
            raise ValueError(f"Duplicate Patient review field: {field}")
        seen.add(field)
        if field in ("patient_details", "prior_information", "current_transcript"):
            if value != text(case.get(field)):
                raise ValueError(f"Source {field} changed for {active['case_id']}; version cases first")
        else:
            active[field] = value or None
    finish()
    if metadata.get("Dataset version") != version:
        raise ValueError("Workbook/case dataset versions differ")
    if {ref["case_id"] for ref in references} != set(known):
        raise ValueError("Patient review cards do not match the supplied cases")
    for ref in references:
        ref["reviewer_id"] = metadata.get("Reviewer alias") or None
    return references


ATTENTION_ORDER = {"1": 1, "2": 2, "3": 3, "Cannot assess": "Unsure"}


def parse_attention_rows(rows, expected_queues, cases, reviewer_id=None):
    """Read the optional Attention exercise (v2 layout) against the versioned cases and queue groups."""
    expected = {q["queue_id"]: q for q in expected_queues}
    source = {c["case_id"]: c for c in cases}
    queues, active, case_id, notes = [], None, None, []

    def finish():
        if active is None:
            return
        if set(active["order"]) != set(expected[active["queue_id"]]["case_ids"]):
            raise ValueError(f"Incomplete attention group {active['queue_id']}")
        active["explanation"] = "\n".join(notes) or None
        filled = all(v is not None for v in active["order"].values())
        active["status"] = "Done" if filled and active["explanation"] else "Not started"

    for row in rows:
        label, value = text(row[0] if row else None), text(row[1] if len(row) > 1 else None)
        group = re.fullmatch(r"Group \d+ / (A\d+)", label)
        if group:
            finish()
            queue_id, case_id, notes = group.group(1), None, []
            if queue_id not in expected or any(q["queue_id"] == queue_id for q in queues):
                raise ValueError(f"Unknown/duplicate attention group: {queue_id}")
            active = {"queue_id": queue_id, "order": {}, "explanation": None, "status": "Not started",
                      "reviewer_id": reviewer_id, "rubric_version": None, "approved": False}
            queues.append(active)
        elif active is None:
            continue
        elif re.fullmatch(r"SYN\d+", label):
            case_id = label
            if case_id not in expected[active["queue_id"]]["case_ids"] or case_id in active["order"]:
                raise ValueError(f"Unknown/duplicate patient in attention group: {case_id}")
            if value != f"ETA {source[case_id]['eta_minutes']} minutes at latest report":
                raise ValueError(f"ETA changed for {case_id}; version cases first")
            active["order"][case_id] = None
        elif label in ("Is ranking meaningful here?", "Why / missing context?"):
            if value:
                notes.append(f"{label} {value}")
        elif case_id is None:
            continue
        elif label == "Age / reported sex / weight" or "EMS Field Report" in label:
            field = ("patient_details" if label.startswith("Age") else
                     "current_transcript" if label.startswith("T=0") else "prior_information")
            if value != text(source[case_id][field]):
                raise ValueError(f"Attention source {field} changed for {case_id}; version cases first")
        elif label == "Attention order":
            if value and value not in ATTENTION_ORDER:
                raise ValueError(f"Invalid attention order for {case_id}: {value}")
            active["order"][case_id] = ATTENTION_ORDER.get(value)
        elif label == "Reason for position / tie / limits":
            if value:
                notes.append(f"{case_id}: {value}")
    finish()
    if {q["queue_id"] for q in queues} != set(expected):
        raise ValueError("Attention groups do not match the queue definitions")
    return queues


def approve_complete(reference):
    """Mark every complete, Done answer approved under the dataset version; return what stayed unapproved."""
    from .contracts import reference_errors
    rubric = reference["dataset_version"]
    skipped = []
    for row in reference["cases"]:
        row["rubric_version"] = rubric
        row["approved"] = True
        errors = reference_errors(row)
        if errors:
            row["approved"] = False
            skipped.append(f"{row['case_id']}: {', '.join(errors)}")
    for queue in reference["queues"]:
        queue["rubric_version"] = rubric
        complete = (queue["status"] == "Done" and queue.get("reviewer_id") and queue.get("explanation")
                    and all(v is not None for v in queue["order"].values()))
        queue["approved"] = bool(complete)
        if not complete:
            skipped.append(f"{queue['queue_id']}: incomplete attention group")
    return skipped


def parse_doctor_rows(rows, cases):
    known = {c["case_id"]: c for c in cases}
    references, active, seen_fields = [], None, set()
    labels = {"1. Brief ED handoff": "reference_summary",
              "2. Important change?": "meaningful_change",
              "Why? What is missing?": "explanation_missing_information",
              "3. Preparation needed?": "preparation_category",
              "Review status": "status"}
    source_labels = {"Earlier report": "prior_information", "Current report": "current_transcript"}

    def finish():
        if active is not None and seen_fields != set(labels.values()) | set(source_labels.values()):
            raise ValueError(f"Incomplete/changed form layout for {active['case_id']}")

    for row in rows:
        label, value = text(row[0] if row else None), row[1] if len(row) > 1 else None
        match = re.search(r"\b(SYN\d+)\b", label) if label.startswith("Case ") else None
        if match:
            finish()
            case_id = match.group(1)
            if case_id not in known or any(r["case_id"] == case_id for r in references):
                raise ValueError(f"Unknown/duplicate case card: {case_id}")
            active, seen_fields = blank_reference(case_id), set()
            references.append(active)
            continue
        if active is None:
            continue
        first_line = label.split("\n")[0]
        field = labels.get(first_line) or source_labels.get(first_line)
        if field:
            if field in seen_fields:
                raise ValueError(f"Duplicate form field: {field}")
            seen_fields.add(field)
            if field in source_labels.values():
                if text(value) != text(known[active["case_id"]][field]):
                    raise ValueError(f"Source report changed for {active['case_id']}; version cases first")
            else:
                active[field] = text(value) or None
    finish()
    if {r["case_id"] for r in references} != set(known):
        raise ValueError("Doctor review cards do not match the supplied cases")
    for ref in references:
        for field, choices in (("meaningful_change", CHANGE), ("preparation_category", PREPARATION),
                               ("status", STATUSES)):
            if ref[field] is not None and ref[field] not in choices:
                raise ValueError(f"Invalid {field} for {ref['case_id']}: {ref[field]}")
    return references


def parse_queue_rows(rows, expected_queues, cases):
    expected = {q["queue_id"]: q for q in expected_queues}
    source = {c["case_id"]: c for c in cases}
    queues, active, case_id, seen = [], None, None, set()

    def finish_patient():
        if case_id is not None and seen != {"Earlier report", "Current report", "Preparation order"}:
            raise ValueError(f"Incomplete/changed queue patient card: {case_id}")

    for row in rows:
        label, value = text(row[0] if row else None), row[1] if len(row) > 1 else None
        group = re.match(r"Group (\d+)\b", label)
        if group:
            finish_patient()
            case_id, seen = None, set()
            queue_id = f"Q{int(group.group(1)):02d}"
            if queue_id not in expected or any(q["queue_id"] == queue_id for q in queues):
                raise ValueError(f"Unknown/duplicate queue: {queue_id}")
            if label != expected[queue_id]["form_heading"]:
                raise ValueError(f"Resource context changed for {queue_id}; version queues first")
            active = {"queue_id": queue_id, "order": {}, "explanation": None,
                      "status": "Not started", "reviewer_id": None, "rubric_version": None,
                      "approved": False}
            queues.append(active)
        elif active:
            match = re.match(r"(SYN\d+)\b", label)
            if match:
                finish_patient()
                case_id = match.group(1)
                seen = set()
                if case_id not in expected[active["queue_id"]]["case_ids"] or case_id in active["order"]:
                    raise ValueError("Unknown/duplicate patient in queue form")
                if label != f"{case_id} · ETA {source[case_id]['eta_minutes']} minutes":
                    raise ValueError(f"ETA changed for {case_id}; version cases first")
                active["order"][case_id] = None
            elif label in ("Earlier report", "Current report", "Preparation order"):
                if case_id is None or label in seen:
                    raise ValueError("Missing patient or duplicate field in queue form")
                seen.add(label)
                if label != "Preparation order":
                    field = "prior_information" if label == "Earlier report" else "current_transcript"
                    if text(value) != text(source[case_id][field]):
                        raise ValueError(f"Queue source report changed for {case_id}; version cases first")
                    continue
                val = text(value)
                if val not in ("", "1", "2", "3", "Unsure"):
                    raise ValueError(f"Invalid rank: {val}")
                active["order"][case_id] = int(val) if val in ("1", "2", "3") else val or None
            elif label.split("\n")[0] == "Why this order?":
                active["explanation"] = text(value) or None
            elif label == "Review status":
                if text(value) not in STATUSES:
                    raise ValueError("Invalid queue status")
                active["status"] = text(value)
    finish_patient()
    if {q["queue_id"] for q in queues} != set(expected):
        raise ValueError("Queue groups do not match expected queues")
    for q in queues:
        if set(q["order"]) != set(expected[q["queue_id"]]["case_ids"]):
            raise ValueError("Incomplete queue layout")
    return queues


def import_workbook(path, cases, queue_definitions):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX import needs openpyxl: pip install -r evals/requirements.txt") from exc
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        patient_review = "Patient review" in workbook.sheetnames
        if patient_review and "Doctor review" in workbook.sheetnames:
            raise ValueError("Workbook contains two review forms; select one source form")
        required = (("Patient review",) + (("Attention exercise",) if queue_definitions else ())) \
            if patient_review else ("Doctor review", "Queue exercise")
        if any(name not in workbook.sheetnames for name in required):
            raise ValueError("Missing required review tabs: " + ", ".join(required))
        grids = {}
        for name in required:
            sheet = workbook[name]
            if (sheet.max_row or 0) > 10000:
                raise ValueError("Unexpected form size")
            grids[name] = []
            for cells in sheet.iter_rows(min_col=1, max_col=2):
                if len(grids[name]) >= 10000:
                    raise ValueError("Unexpected form size")
                # The progress counter is the only expected formula. Never evaluate formulas.
                if any(c.data_type == "f" for c in cells) and not (name == "Doctor review" and cells[0].row == 5):
                    raise ValueError(f"Formula in source/answer area: {name}, row {cells[0].row}")
                grids[name].append([c.value for c in cells])
        references = (parse_patient_rows(grids["Patient review"], cases) if patient_review
                      else parse_doctor_rows(grids["Doctor review"], cases))
        if "Attention exercise" in grids:
            reviewer = references[0]["reviewer_id"] if references else None
            queues = parse_attention_rows(grids["Attention exercise"], queue_definitions, cases, reviewer)
        else:
            queues = parse_queue_rows(grids["Queue exercise"], queue_definitions, cases) \
                if "Queue exercise" in grids else []
        return {"kind": "clinician_reference", "dataset_version": case_dataset_version(cases) or "pilot_v1",
                "source": {"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()},
                "cases": references, "queues": queues}
    finally:
        workbook.close()
