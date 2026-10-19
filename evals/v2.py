"""Offline v2 intake/scoring; no provider imports, schema crosswalks or approvals.

The unchanged blank form is the structural reference. Static source edits are
retained as blocking issues, not silently accepted as a new patient dataset.
"""

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

from .contracts import CHANGE, ROOT, load_cases, load_json, number, text
from .metrics import classification_metrics, pairwise_metrics, percentile, ratio
from .report import save_json, table, write_csv

FORM = "clinician_form_v2"
DATASET = "pilot_v1"
RUBRIC = "review_attention_v2_draft"
EXERCISES = "attention_v2"
TEMPLATE = ROOT / "docs/clinician-annotation/v2/Annotation_Notebook_v2.xlsx"
ATTENTION = ("Immediate review", "Review soon", "Routine awareness", "Cannot assess")
DIRECTION = ("Improved", "Worsened", "Mixed", "Other", "Unclear")
MEANINGFUL = ("Yes", "Only some comparisons", "No", "Cannot assess")
CASE_LABELS = {"EMS to ED Handoff": "summary", "Meaningful change?": "meaningful_change",
               "Direction (if Yes)": "change_direction", "Evidence / missing information": "evidence_missing_information",
               "Attention before arrival": "attention_category", "Reason for attention / limits": "attention_reason"}
CORE = tuple(k for k in CASE_LABELS.values() if k != "change_direction")
GROUPS = {"A01": ["SYN001", "SYN004", "SYN007"], "A02": ["SYN002", "SYN008", "SYN010"]}
VERSIONS = {"form_version": FORM, "dataset_version": DATASET, "rubric_version": RUBRIC,
            "exercise_version": EXERCISES}


def read_grids(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        grids = {}
        for sheet in workbook:
            key = sheet.title.casefold()
            if key in grids:
                raise ValueError("Duplicate case-insensitive tab name")
            if (sheet.max_row or 0) > 1000 or (sheet.max_column or 0) > 32:
                raise ValueError("Unexpected v2 workbook size")
            rows = []
            for cells in sheet.iter_rows():
                if len(rows) >= 1000 or len(cells) > 32:
                    raise ValueError("Unexpected v2 workbook size")
                if any(c.data_type in ("f", "e") for c in cells):
                    raise ValueError(f"Formula/error cell in {sheet.title}; do not evaluate annotation formulas")
                rows.append([c.value for c in cells])
            grids[key] = rows
        required = {"patient review", "attention exercise", "form feedback"}
        if set(grids) != required:
            raise ValueError("Expected only Patient review, Attention exercise, Form feedback tabs")
        return grids
    finally:
        workbook.close()


def cell(rows, row, column=1):
    return rows[row - 1][column] if len(rows) >= row and len(rows[row - 1]) > column else None


def issue(code, sheet, row, expected=None, observed=None):
    return {"code": code, "sheet": sheet, "cell": f"B{row}", "expected": expected, "observed": observed}


def content_errors(ref):
    errors = [f"missing {k}" for k in CORE if not text(ref.get(k))]
    for key in ("summary", "evidence_missing_information", "attention_reason"):
        if ref.get(key) is not None and not isinstance(ref[key], str):
            errors.append(f"invalid narrative {key}")
    for key, choices in (("meaningful_change", CHANGE), ("attention_category", ATTENTION)):
        if ref.get(key) is not None and ref[key] not in choices:
            errors.append(f"invalid {key}")
    if ref.get("meaningful_change") == "Yes" and ref.get("change_direction") not in DIRECTION:
        errors.append("missing/invalid change_direction for Yes")
    elif ref.get("change_direction") is not None and ref["change_direction"] not in DIRECTION:
        errors.append("invalid change_direction")
    return errors + [x["code"] for x in ref.get("issues", [])]


def approval_errors(ref):
    errors = []
    if ref.get("approved") is not True:
        errors.append("not approved")
    for key in ("reviewer_id", "rubric_version", "reviewed_at", "approval_note"):
        if not text(ref.get(key)):
            errors.append(f"missing {key}")
    return errors


def parse_grids(grids, template, cases):
    """Fixed v2 layout: inserted/missing question rows require an explicit version change.

    Extra feedback cells are preserved. Tab-name capitalization is insignificant.
    No clinical free text is normalized or converted to another category.
    """
    known = {c["case_id"]: c for c in cases}
    for name in ("patient review", "attention exercise"):
        cid = None
        for r, row in enumerate(template[name], 1):
            label = text(row[0] if row else None)
            match = re.fullmatch(r"Case \d+ / (SYN\d+)", label)
            if match:
                cid = match[1]
            elif label in known:
                cid = label
            if cid and cid not in known:
                raise ValueError("Blank form references unknown source case")
            if cid and "EMS Field Report" in label:
                field = "current_transcript" if label.startswith("T=0\n") else "prior_information"
                expected_label = ("T=0\n" + ("EMS Field Report" if known[cid]["update_number"] == 1 else "Updated EMS Field Report")) if field == "current_transcript" else f"T-{known[cid]['elapsed_minutes']} min\nEMS Field Report"
                if label != expected_label or text(cell(template[name], r)) != text(known[cid][field]):
                    raise ValueError("Blank v2 template source differs from canonical case/time")
            if cid and (match or label in known):
                if cell(template[name], r) != f"ETA {known[cid]['eta_minutes']} minutes at latest report":
                    raise ValueError("Blank form ETA differs from canonical source")
    cases_out, groups_out, global_issues = [], [], []
    ownership, editable = {}, set()
    approvals = {"reviewer_id": None, "rubric_version": None, "reviewed_at": None,
                 "approval_note": None, "approved": False}
    current, group, member = None, None, None
    for name in ("patient review", "attention exercise"):
        current = group = member = None
        for r, row in enumerate(template[name], 1):
            label = text(row[0] if row else None)
            case_match = re.fullmatch(r"Case \d+ / (SYN\d+)", label)
            group_match = re.fullmatch(r"Group \d+ / (A\d+)", label)
            if name == "patient review" and case_match:
                cid = case_match[1]
                if cid not in known:
                    raise ValueError("V2 template and source cases differ; version them together")
                current = {"case_id": cid, **dict.fromkeys(CASE_LABELS.values()), "issues": [], **approvals}
                cases_out.append(current)
            if name == "attention exercise" and group_match:
                gid = group_match[1]
                group = {"exercise_id": gid, "ranking_meaningful": None, "explanation": None,
                         "order": {}, "reasons": {}, "issues": [], **approvals}
                groups_out.append(group)
            if name == "attention exercise" and label in known:
                member = label
            owner = current if name == "patient review" else group
            if owner is not None:
                ownership[name, r] = owner
            value = cell(grids[name], r)
            if name == "patient review" and label in ("Reviewer alias", "Role / relevant experience"):
                editable.add((name, r))
            if name == "patient review" and current is not None and label in CASE_LABELS:
                current[CASE_LABELS[label]] = value if text(value) else None
                editable.add((name, r))
            if name == "attention exercise" and group is not None:
                key = {"Is ranking meaningful here?": "ranking_meaningful", "Why / missing context?": "explanation"}.get(label)
                if key:
                    group[key] = value if text(value) else None
                    editable.add((name, r))
                if label in ("Attention order", "Reason for position / tie / limits"):
                    if member not in GROUPS.get(group["exercise_id"], []):
                        raise ValueError("Unknown attention-exercise member")
                    key = "order" if label == "Attention order" else "reasons"
                    group[key][member] = value if text(value) else None
                    editable.add((name, r))
    if {r["case_id"] for r in cases_out} != set(known) or len(cases_out) != len(known):
        raise ValueError("V2 template does not cover exactly the source cases")
    if {g["exercise_id"] for g in groups_out} != set(GROUPS):
        raise ValueError("Unexpected v2 attention exercises")

    # Every static label/value must match the frozen blank form. The canonical
    # case text additionally cross-checks the blank form, not only the upload.
    for name in ("patient review", "attention exercise", "form feedback"):
        for r, expected_row in enumerate(template[name], 1):
            label = text(expected_row[0] if expected_row else None)
            observed_label = text(cell(grids[name], r, 0))
            if label != observed_label:
                raise ValueError(f"V2 layout changed: {name}!A{r}; version form first")
            if name == "form feedback" and r in (5, 6, 7, 8, 9, 10):
                continue
            if (name, r) in editable:
                continue
            expected, observed = cell(template[name], r), cell(grids[name], r)
            if text(expected) != text(observed):
                found = issue("static_source_or_instruction_changed", name, r, expected, observed)
                owner = ownership.get((name, r))
                (owner["issues"] if owner is not None else global_issues).append(found)
        if name != "form feedback":
            for r in range(len(template[name]) + 1, len(grids[name]) + 1):
                if any(text(v) for v in grids[name][r - 1]):
                    global_issues.append(issue("unexpected_extra_case_rows", name, r))

    for ref in cases_out:
        ref["filled_fields"] = sum(text(ref.get(k)) != "" for k in CASE_LABELS.values())
        ref["content_errors"] = content_errors(ref)
        ref["intake_status"] = "not_started" if not ref["filled_fields"] else "needs_review" if ref["content_errors"] else "filled_pending_review"
    for g in groups_out:
        if set(g["order"]) != set(GROUPS[g["exercise_id"]]):
            raise ValueError("Attention-exercise layout has missing members")
    return {"kind": "clinician_reference_v2", **VERSIONS, "rubric_frozen": False,
            "reviewer_metadata": {"alias": cell(grids["patient review"], 3), "role": cell(grids["patient review"], 4)},
            "global_issues": global_issues, "cases": cases_out, "exercises": groups_out,
            "form_feedback": [{"row": i, "values": row} for i, row in enumerate(grids["form feedback"], 1)
                              if i in (5, 6, 7, 8, 9, 10) or i > len(template["form feedback"]) if any(text(v) for v in row)],
            "raw_nonempty_cells": {name: [{"row": i, "values": row} for i, row in enumerate(rows, 1)
                                          if any(text(v) for v in row)] for name, rows in grids.items()}}


def import_v2(path, cases):
    result = parse_grids(read_grids(path), read_grids(TEMPLATE), cases)
    result["source"] = {"file": Path(path).name, "sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                        "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()}
    return result


def baseline_v2(cases):
    """Raw-source echo plus explicit abstention; never uses reference answers."""
    return {"kind": "baseline_predictions_v2", **VERSIONS,
            "config": {"config_id": "raw_report_abstention", "provider": "deterministic", "model": "none"},
            "records": [{"case_id": c["case_id"], "status": "ok", "output": {
                "summary": c["current_transcript"], "meaningful_change": "Unclear", "change_direction": None,
                "evidence_missing_information": "Raw-report baseline; no clinical interpretation performed.",
                "attention_category": "Cannot assess", "attention_reason": "Baseline abstains from attention judgments."}}
                        for c in cases], "exercises": []}


def output_errors(output):
    if not isinstance(output, dict) or set(output) != set(CASE_LABELS.values()):
        return ["output fields do not match v2 contract"]
    errors = content_errors(output)
    if output.get("meaningful_change") != "Yes" and output.get("change_direction") is not None:
        errors.append("change_direction must be null unless meaningful_change is Yes")
    return errors


def rank(value):
    if value == "Cannot assess":
        return value
    if isinstance(value, bool):
        return None
    if number(value, integer=True, minimum=1) and value <= 3:
        return int(value)
    if isinstance(value, str) and value in ("1", "2", "3"):
        return int(value)
    return None


def score_v2(cases, references, predictions, *, exercise_ids=None):
    if references.get("kind") != "clinician_reference_v2" or predictions.get("kind") not in ("baseline_predictions_v2", "model_predictions_v2"):
        raise ValueError("V2 references and v2 predictions required; mock/v1 mixing is not allowed")
    for key, value in VERSIONS.items():
        if references.get(key) != value or predictions.get(key) != value:
            raise ValueError(f"Mismatched v2 {key}; do not crosswalk versions")
    config = predictions.get("config", {})
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]*", text(config.get("config_id"))):
        raise ValueError("Invalid v2 config_id")
    if any(not isinstance(config.get(k), str) or not text(config[k]) for k in ("provider", "model")):
        raise ValueError("V2 config requires provider and model provenance")
    known = {c["case_id"] for c in cases}
    if exercise_ids is None:
        exercise_ids = list(GROUPS)
    if (not isinstance(exercise_ids, list) or len(set(exercise_ids)) != len(exercise_ids)
            or set(exercise_ids) - set(GROUPS)):
        raise ValueError("Unknown/duplicate v2 attention exercise selection")
    groups = {gid: GROUPS[gid] for gid in exercise_ids}
    if exercise_ids and any(not set(ids) <= known for ids in groups.values()):
        raise ValueError("Selected attention exercise contains cases outside evaluation subset")

    def index(rows, field, allowed):
        out = {}
        for row in rows:
            key = row[field]
            if key not in allowed or key in out:
                raise ValueError("Duplicate/unknown v2 record")
            out[key] = row
        return out

    refs = index(references.get("cases", []), "case_id", known)
    records = index(predictions.get("records", []), "case_id", known)
    valid = {}
    for cid, rec in records.items():
        if rec.get("status") not in ("ok", "error", "timeout", "refusal"):
            raise ValueError("Invalid v2 prediction status")
        for key in ("latency_ms", "cost_usd", "input_tokens", "output_tokens", "retry_count"):
            if rec.get(key) is not None and not number(rec[key], integer=key in ("input_tokens", "output_tokens", "retry_count")):
                raise ValueError(f"Invalid v2 telemetry: {key}")
        valid[cid] = rec["output"] if rec["status"] == "ok" and not output_errors(rec.get("output")) else {}
    eligible, excluded = {}, {}
    for cid in sorted(known):
        ref = refs.get(cid)
        errors = ["missing reference"] if ref is None else content_errors(ref) + approval_errors(ref)
        if references.get("global_issues"):
            errors.append("global form/source issues unresolved")
        if references.get("rubric_frozen") is not True:
            errors.append("rubric not frozen")
        if ref and ref.get("rubric_version") != references["rubric_version"]:
            errors.append("record rubric version differs")
        if errors:
            excluded[cid] = errors
        else:
            eligible[cid] = ref
    attention = classification_metrics([r["attention_category"] for r in eligible.values()],
        [valid.get(cid, {}).get("attention_category", "[unavailable]") for cid in eligible], ATTENTION)
    change_refs = {cid: r for cid, r in eligible.items() if r["meaningful_change"] in ("Yes", "No")}
    change = classification_metrics([r["meaningful_change"] for r in change_refs.values()],
        [valid.get(cid, {}).get("meaningful_change", "[unavailable]") for cid in change_refs], ("Yes", "No"))
    direction_refs = {cid: r for cid, r in eligible.items() if r["meaningful_change"] == "Yes"}
    direction = classification_metrics([r["change_direction"] for r in direction_refs.values()],
        [valid.get(cid, {}).get("change_direction") or "[unavailable]" for cid in direction_refs], DIRECTION)
    ref_groups = index(references.get("exercises", []), "exercise_id", set(groups))
    pred_groups = index(predictions.get("exercises", []), "exercise_id", set(groups))
    exercise_results = []
    for gid, ids in groups.items():
        ref = ref_groups.get(gid, {})
        reasons = approval_errors(ref) + [x["code"] for x in ref.get("issues", [])]
        if references.get("global_issues"):
            reasons.append("global form/source issues unresolved")
        if references.get("rubric_frozen") is not True:
            reasons.append("rubric not frozen")
        if ref.get("rubric_version") != references["rubric_version"]:
            reasons.append("record rubric version differs")
        if ref.get("ranking_meaningful") not in ("Yes", "Only some comparisons"):
            reasons.append("ranking not established as meaningful")
        source_order = ref.get("order") if isinstance(ref.get("order"), dict) else {}
        order = {cid: rank(source_order.get(cid)) for cid in ids}
        if set(source_order) != set(ids) or any(v is None for v in order.values()):
            reasons.append("incomplete/invalid ranks")
        if not text(ref.get("explanation")) or any(not text(ref.get("reasons", {}).get(cid)) for cid in ids):
            reasons.append("missing ranking explanation")
        candidate_record = pred_groups.get(gid, {})
        candidate = candidate_record.get("order") if isinstance(candidate_record.get("order"), dict) and candidate_record.get("status", "ok") == "ok" else {}
        predicted = {cid: rank(candidate.get(cid)) for cid in ids} if set(candidate) == set(ids) else {}
        metric = pairwise_metrics(order, predicted) if not reasons else pairwise_metrics({}, {})
        exercise_results.append({"exercise_id": gid, "eligible": not reasons, "exclusion_reasons": reasons, **metric})
    intake = [{"case_id": cid, "filled_fields": sum(bool(text(refs.get(cid, {}).get(k))) for k in CASE_LABELS.values()),
               "core_fields_complete": not any(not text(refs.get(cid, {}).get(k)) for k in CORE),
               "blocking_issue_count": len(content_errors(refs.get(cid, {}))),
               "reference_eligible": cid in eligible, "prediction_usable": bool(valid.get(cid))} for cid in sorted(known)]
    latencies = [r["latency_ms"] for r in records.values() if r.get("latency_ms") is not None]
    costs = [r["cost_usd"] for r in records.values() if r.get("cost_usd") is not None]
    usable = sum(bool(valid.get(cid)) for cid in known)
    checked = sum(r.get("provider_status", r["status"]) == "ok" for r in records.values())
    statuses = Counter("missing" if cid not in records else "invalid" if records[cid]["status"] == "ok" and not valid.get(cid) else records[cid]["status"] for cid in known)
    return {"kind": "v2_evaluation", **VERSIONS,
            "mode": "OFFLINE V2 PIPELINE CHECK - deterministic baseline, not LLM performance" if predictions["kind"] == "baseline_predictions_v2" else "V2 precomputed model evaluation",
            "config": config, "cases_total": len(known), "cases_started": sum(r["filled_fields"] > 0 for r in intake),
            "approved_reference_cases": len(eligible), "excluded_cases": excluded, "case_intake": intake,
            "source_issues": [{"scope": cid, **item} for cid, ref in refs.items() for item in ref.get("issues", [])]
                             + [{"scope": "form", **item} for item in references.get("global_issues", [])],
            "usable_response_rate": ratio(usable, len(known)), "error_rate": ratio(len(known) - usable, len(known)),
            "schema_validation_rate": ratio(usable, checked), "expected_responses": len(known),
            "returned_responses": len(records), "status_counts": dict(statuses),
            "attention": attention, "change": change, "direction": direction, "exercises": exercise_results,
            "case_predictions": [{"case_id": cid, "status": records.get(cid, {}).get("status", "missing"),
                                  "output": records.get(cid, {}).get("output"),
                                  "schema_errors": records.get(cid, {}).get("schema_errors", [])} for cid in sorted(known)],
            "latency_median_ms": percentile(latencies, .5), "latency_p95_ms": percentile(latencies, .95),
            "latency_coverage_cases": len(latencies), "reported_case_cost_usd": sum(costs) if costs else None,
            "case_cost_coverage": ratio(len(costs), len(known)),
            "summary_factual_precision": None, "summary_content_recall": None,
            "summary_note": "Separate human summary grading is required; no automatic textual similarity score substituted.",
            "latency_ms": None, "cost_usd": 0 if predictions["kind"] == "baseline_predictions_v2" else sum(costs) if len(costs) == len(known) else None,
            "model_calls": 0 if predictions["kind"] == "baseline_predictions_v2" else None}


def save_v2_report(report, out):
    import html
    out = Path(out)
    out.mkdir(parents=True, exist_ok=False)
    save_json(out / "metrics.json", report)
    write_csv(out / "case_results.csv", report["case_intake"])
    comparison = {"mode": report["mode"], "config_id": report["config"]["config_id"],
                  "cases_total": report["cases_total"], "cases_started": report["cases_started"],
                  "approved_reference_cases": report["approved_reference_cases"],
                  "attention_accuracy": report["attention"]["accuracy"], "attention_macro_f1": report["attention"]["macro_f1"],
                  "immediate_review_false_negative_rate": report["attention"]["per_class"]["Immediate review"]["false_negative_rate"],
                  "change_recall": report["change"]["per_class"]["Yes"]["recall"], "change_f1": report["change"]["per_class"]["Yes"]["f1"],
                  "direction_accuracy": report["direction"]["accuracy"], "usable_response_rate": report["usable_response_rate"],
                  "schema_validation_rate": report["schema_validation_rate"], "error_rate": report["error_rate"],
                  "expected_responses": report["expected_responses"], "returned_responses": report["returned_responses"],
                  "latency_median_ms": report["latency_median_ms"], "latency_p95_ms": report["latency_p95_ms"],
                  "latency_coverage_cases": report["latency_coverage_cases"], "reported_case_cost_usd": report["reported_case_cost_usd"],
                  "case_cost_coverage": report["case_cost_coverage"],
                  "summary_factual_precision": None, "summary_content_recall": None, "model_calls": report["model_calls"]}
    write_csv(out / "comparison.csv", [comparison])
    write_csv(out / "exercise_results.csv", [{k: v for k, v in r.items() if k != "exclusion_reasons"} for r in report["exercises"]])
    body = ["<!doctype html><html><head><meta charset='utf-8'><title>EMS v2 pipeline check</title>",
            "<style>body{font:16px system-ui;max-width:1100px;margin:32px}table{border-collapse:collapse;display:block;overflow:auto}th,td{padding:9px;border:1px solid #ccc;text-align:left}.notice{background:#fff2cc;padding:16px}h2{margin-top:32px}</style></head><body>",
            "<h1>EMS v2 evaluation</h1><p class='notice'>", html.escape(report["mode"]),
            "<br>Blank/N/A means unavailable, not zero. Source edits and unapproved answers cannot become clinical reference scores.</p>",
            f"<p>Cases started: {report['cases_started']} / {report['cases_total']}. Approved references: {report['approved_reference_cases']}.</p>",
            "<h2>Scoring output</h2>", table(["Metric", "Value"], list(comparison.items())),
            "<h2>Annotation coverage</h2>", table(list(report["case_intake"][0]), [list(r.values()) for r in report["case_intake"]]),
            "<h2>Collected outputs (not human-graded summaries)</h2>",
            table(["Case", "Status", "Output", "Schema errors"], [[r["case_id"], r["status"], r["output"], r["schema_errors"]] for r in report["case_predictions"]]),
            "<h2>Why references are excluded</h2>", table(["Case", "Reasons"], [[cid, "; ".join(errors)] for cid, errors in report["excluded_cases"].items()]),
            "<h2>Source/form edits requiring review</h2>", table(["Scope", "Sheet", "Cell", "Expected", "Uploaded"], [[i["scope"], i["sheet"], i["cell"], i["expected"], i["observed"]] for i in report["source_issues"]]),
            "<h2>Attention exercises</h2>", table(["Exercise", "Eligible", "Pairwise accuracy", "Reasons"], [[r["exercise_id"], r["eligible"], r["pairwise_accuracy"], "; ".join(r["exclusion_reasons"])] for r in report["exercises"]]),
            "<p>Summary quality, clinical safety, actual lead time, and product utility are not established by this software check.</p></body></html>"]
    (out / "dashboard.html").write_text("".join(body), encoding="utf-8")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Offline v2 annotation intake and scoring; never calls models")
    parser.add_argument("command", choices=("import", "smoke", "score"))
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--references", type=Path)
    parser.add_argument("--predictions", type=Path)
    parser.add_argument("--out", type=Path, required=True, help="New directory; never overwritten")
    args = parser.parse_args(argv)
    try:
        if args.out.exists():
            raise ValueError("Output directory already exists; choose a new run directory")
        cases = load_cases()
        if args.command in ("import", "smoke"):
            if args.workbook is None:
                raise ValueError("--workbook required")
            refs = import_v2(args.workbook, cases)
            if args.command == "import":
                args.out.mkdir(parents=True)
                save_json(args.out / "references.json", refs)
            else:
                preds = baseline_v2(cases)
                report = score_v2(cases, refs, preds)
                save_v2_report(report, args.out)
                save_json(args.out / "references.json", refs)
                save_json(args.out / "baseline_predictions.json", preds)
                print(f"Cases started: {report['cases_started']}/{report['cases_total']}; approved: {report['approved_reference_cases']}.")
                print("OFFLINE CHECK ONLY: raw-report baseline, zero model calls. Unapproved quality metrics are N/A.")
        else:
            if args.references is None or args.predictions is None:
                raise ValueError("--references and --predictions required")
            report = score_v2(cases, load_json(args.references), load_json(args.predictions))
            save_v2_report(report, args.out)
        print(f"Created {args.out.resolve()}")
        return 0
    except (ValueError, KeyError, TypeError, OSError, RuntimeError, ImportError) as exc:
        print(f"V2 evaluation error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
