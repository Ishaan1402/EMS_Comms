import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from evals.annotations import import_workbook, parse_patient_rows
from evals.contracts import ROOT, case_dataset_version, load_cases, load_json, reference_errors
from evals.demo import raw_report_baseline
from evals.evaluate import evaluate
from evals.experiments import build_plan, case_input


CASES = ROOT / "docs/clinician-annotation/patient_review_v2_1_cases.json"


def form_rows(cases):
    rows = [["Dataset version", "patient_review_v2.1"]]
    for index, case in enumerate(cases, 1):
        rows += [[f"Case {index} / {case['case_id']}", f"ETA {case['eta_minutes']} minutes at latest report"],
                 ["Age / reported sex / weight", case["patient_details"]]]
        if case["prior_information"] != "No earlier report provided.":
            rows.append([f"T-{case['elapsed_minutes']} min\nEMS Field Report", case["prior_information"]])
        rows += [["T=0\nUpdated EMS Field Report", case["current_transcript"]],
                 ["EMS to ED Handoff", ""], ["Meaningful change?", ""], ["Direction (if Yes)", ""],
                 ["Evidence / missing information", ""], ["Preparation needed?", ""],
                 ["Reason for preparation / limits", ""], ["Review status", "Not started"], ["", ""]]
    return rows


class PatientReviewTests(unittest.TestCase):
    def setUp(self):
        self.cases = load_cases(CASES)
        self.rows = form_rows(self.cases)

    def set_answer(self, label, value):
        next(row for row in self.rows if row[0] == label)[1] = value

    def use_esi_form(self):
        labels = {"Preparation needed?": "ESI level (1–5)",
                  "Reason for preparation / limits": "ESI reason / limits"}
        for row in self.rows:
            row[0] = labels.get(row[0], row[0])

    def test_native_esi_levels_remain_distinct_and_keep_legacy_metric_mapping(self):
        self.use_esi_form()
        for level, category in ((1, "Prepare now"), (2, "Prepare now"), (3, "Can wait"),
                                (4, "Routine"), (5, "Routine"), ("Cannot assess", "Unsure")):
            with self.subTest(level=level):
                self.set_answer("ESI level (1–5)", level)
                self.set_answer("ESI reason / limits", "Reported findings.\nImportant limits remain.")
                ref = parse_patient_rows(self.rows, self.cases)[0]
                self.assertEqual(ref["esi_level"], str(level))
                self.assertEqual(ref["esi_reason"], "Reported findings.\nImportant limits remain.")
                self.assertEqual(ref["preparation_category"], category)
                self.assertIn("ESI / limits: Reported findings.", ref["explanation_missing_information"])
                self.assertFalse(ref["approved"])

    def test_blank_esi_is_missing_rather_than_a_default_or_abstention(self):
        self.use_esi_form()
        ref = parse_patient_rows(self.rows, self.cases)[0]
        self.assertIsNone(ref["esi_level"])
        self.assertIsNone(ref["preparation_category"])
        self.assertIsNone(ref["preparation_category_display"])
        self.assertFalse(ref["approved"])

    def test_integral_excel_scores_and_ascii_esi_label_import(self):
        self.use_esi_form()
        row = next(row for row in self.rows if row[0] == "ESI level (1–5)")
        row[:] = ["ESI level (1-5)", 2.0]
        self.rows.insert(6, ["Risk Score (1-10)", 7.0])
        self.rows.insert(7, ["Risk reason / limits", "Reported deterioration."])
        ref = parse_patient_rows(self.rows, self.cases)[0]
        self.assertEqual(ref["esi_level"], "2")
        self.assertEqual(ref["risk_score"], 7)
        self.assertEqual(ref["risk_reason"], "Reported deterioration.")
        self.assertNotIn("risk_score", parse_patient_rows(self.rows, self.cases)[1])

    def test_invalid_esi_is_rejected_instead_of_clamped_or_mapped(self):
        self.use_esi_form()
        for level in (0, 6, 1.5, True, "ESI 2", "Unsure", "nan"):
            with self.subTest(level=level):
                self.set_answer("ESI level (1–5)", level)
                with self.assertRaisesRegex(ValueError, "Invalid esi_level"):
                    parse_patient_rows(self.rows, self.cases)

    def test_optional_risk_score_does_not_accept_invalid_or_fractional_values(self):
        self.use_esi_form()
        row = ["Risk Score (1-10)", ""]
        self.rows.insert(6, row)
        for value in (0, 11, 4.5, True, "7/10"):
            with self.subTest(value=value):
                row[1] = value
                with self.assertRaisesRegex(ValueError, "Invalid risk_score"):
                    parse_patient_rows(self.rows, self.cases)
        row[1] = ""
        self.assertIsNone(parse_patient_rows(self.rows, self.cases)[0]["risk_score"])

    def test_esi_form_rejects_duplicate_aliases_mixed_formats_and_missing_reason_row(self):
        for extra in (["ESI level", "2"], ["Preparation needed?", "Can wait"]):
            with self.subTest(extra=extra):
                self.rows = form_rows(self.cases)
                self.use_esi_form()
                self.rows.insert(6, extra)
                with self.assertRaisesRegex(ValueError, "Duplicate Patient review field|Mixed ESI"):
                    parse_patient_rows(self.rows, self.cases)
        self.rows = form_rows(self.cases)
        self.use_esi_form()
        self.rows.remove(next(row for row in self.rows if row[0] == "ESI reason / limits"))
        with self.assertRaisesRegex(ValueError, "Incomplete/changed"):
            parse_patient_rows(self.rows, self.cases)

    def test_esi_form_still_rejects_changed_source_reports(self):
        self.use_esi_form()
        next(row for row in self.rows if "324 mg" in row[1])[1] = "A rewritten source report."
        with self.assertRaisesRegex(ValueError, "current_transcript changed"):
            parse_patient_rows(self.rows, self.cases)

    def test_esi_approval_requires_its_reason_and_consistent_derived_category(self):
        from evals.annotations import approve_complete
        self.use_esi_form()
        self.rows.insert(0, ["Reviewer alias", "Dr Test"])
        for label, value in (("EMS to ED Handoff", "Handoff."), ("Meaningful change?", "No earlier report"),
                             ("Evidence / missing information", "Source evidence."),
                             ("ESI level (1–5)", "2"), ("Review status", "Done")):
            self.set_answer(label, value)
        reference = {"dataset_version": "patient_review_v2.1",
                     "cases": parse_patient_rows(self.rows, self.cases), "queues": []}
        self.assertIn("missing esi_reason", approve_complete(reference)[0])
        self.assertFalse(reference["cases"][0]["approved"])
        self.set_answer("ESI reason / limits", "Reported high-risk features.")
        reference["cases"] = parse_patient_rows(self.rows, self.cases)
        approve_complete(reference)
        ref = reference["cases"][0]
        self.assertTrue(ref["approved"])
        ref["preparation_category"] = "Routine"
        self.assertIn("preparation_category does not match esi_level", reference_errors(ref))

    def test_real_xlsx_import_preserves_original_comments_without_approving_them(self):
        try:
            from openpyxl import Workbook
            from openpyxl.comments import Comment
        except ImportError:
            self.skipTest("Optional openpyxl not installed")
        self.use_esi_form()
        self.set_answer("ESI level (1–5)", 3)
        self.set_answer("ESI reason / limits", "Reported resource needs.")
        book = Workbook()
        sheet = book.active
        sheet.title = "Patient review"
        for row in self.rows:
            sheet.append(row)
        handoff_row = next(i for i, row in enumerate(self.rows, 1) if row[0] == "EMS to ED Handoff")
        address = f"B{handoff_row}"
        note = "Assistant-formatted summary.\nOriginal reviewer text: Good information, can structure  "
        sheet[address].comment = Comment(note, "Reviewer")
        archive = book.create_sheet("Original v2")
        archive["B1"] = "Not a current answer."
        archive["B1"].comment = Comment("Archived note must not enter current references.", "Reviewer")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "review.xlsx"
            book.save(path)
            book.close()
            result = import_workbook(path, self.cases, [])
        self.assertEqual(result["cases"][0]["esi_level"], "3")
        self.assertEqual(result["source"]["cell_notes"], {
            "Patient review": {address: {"author": "Reviewer", "text": note}}})
        self.assertTrue(all(row["approved"] is False for row in result["cases"]))

    def test_partial_answers_and_optional_direction_survive_without_approval(self):
        self.set_answer("EMS to ED Handoff", "Partial clinician handoff, still under discussion.")
        self.set_answer("Preparation needed?", "Cannot assess")
        self.set_answer("Evidence / missing information", "Important history is not established.")
        self.set_answer("Reason for preparation / limits", "I cannot select preparation needs yet.")
        self.set_answer("Review status", "Done")
        reference = parse_patient_rows(self.rows, self.cases)[0]
        self.assertEqual(reference["reference_summary"], "Partial clinician handoff, still under discussion.")
        self.assertEqual(reference["preparation_category"], "Unsure")
        self.assertEqual(reference["preparation_category_display"], "Cannot assess")
        self.assertIsNone(reference["change_direction"])
        self.assertIn("Important history is not established.", reference["explanation_missing_information"])
        self.assertIn("I cannot select preparation needs yet.", reference["explanation_missing_information"])
        self.assertFalse(reference["approved"])
        self.assertIsNone(reference["reviewer_id"])

    def test_all_app_display_categories_import(self):
        for choice in ("Prepare now", "Can wait", "Routine", "Cannot assess"):
            with self.subTest(choice=choice):
                self.set_answer("Preparation needed?", choice)
                reference = parse_patient_rows(self.rows, self.cases)[0]
                self.assertEqual(reference["preparation_category"], "Unsure" if choice == "Cannot assess" else choice)
                self.assertEqual(reference["preparation_category_display"], choice)

    def test_review_urgency_is_not_silently_translated_to_preparation(self):
        self.set_answer("Preparation needed?", "Immediate review")
        with self.assertRaisesRegex(ValueError, "Invalid preparation_category"):
            parse_patient_rows(self.rows, self.cases)

    def test_demographic_edit_is_rejected_instead_of_grading_different_inputs(self):
        self.set_answer("Age / reported sex / weight", "5 years / male. Weight 20 kg.")
        with self.assertRaisesRegex(ValueError, "patient_details changed"):
            parse_patient_rows(self.rows, self.cases)

    def test_changed_medication_number_in_report_is_rejected(self):
        row = next(row for row in self.rows if "324 mg" in row[1])
        row[1] = row[1].replace("324 mg", "32 mg")
        with self.assertRaisesRegex(ValueError, "current_transcript changed"):
            parse_patient_rows(self.rows, self.cases)

    def test_missing_earlier_report_is_rejected(self):
        row = next(row for row in self.rows if row[0].startswith("T-8 min"))
        self.rows.remove(row)
        with self.assertRaisesRegex(ValueError, "Incomplete/changed"):
            parse_patient_rows(self.rows, self.cases)

    def test_eta_or_interval_edits_are_rejected(self):
        self.rows[1][1] = "ETA 2 minutes at latest report"
        with self.assertRaisesRegex(ValueError, "ETA changed"):
            parse_patient_rows(self.rows, self.cases)
        self.rows = form_rows(self.cases)
        next(row for row in self.rows if row[0].startswith("T-8 min"))[0] = "T-3 min\nEMS Field Report"
        with self.assertRaisesRegex(ValueError, "interval changed"):
            parse_patient_rows(self.rows, self.cases)

    def test_old_dataset_or_duplicate_answer_cannot_be_imported_as_new(self):
        self.rows[0][1] = "pilot_v1"
        with self.assertRaisesRegex(ValueError, "dataset versions differ"):
            parse_patient_rows(self.rows, self.cases)
        self.rows = form_rows(self.cases)
        self.rows.insert(5, ["Preparation needed?", "Routine"])
        with self.assertRaisesRegex(ValueError, "Duplicate Patient review field"):
            parse_patient_rows(self.rows, self.cases)

    def test_additional_unmatched_patient_information_is_rejected(self):
        self.rows.insert(4, ["Other patient history", "New diagnosis not present in the app input."])
        with self.assertRaisesRegex(ValueError, "Unexpected Patient review field"):
            parse_patient_rows(self.rows, self.cases)

    def test_unversioned_or_mixed_fixtures_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "versioned case fixtures"):
            parse_patient_rows(self.rows, load_cases())
        self.cases[0]["dataset_version"] = "pilot_v1"
        with self.assertRaisesRegex(ValueError, "versions differ"):
            case_dataset_version(self.cases)

    def test_changed_cases_cannot_be_scored_with_old_matching_reference_prediction_versions(self):
        references = {"kind": "clinician_reference", "dataset_version": "pilot_v1", "cases": [], "queues": []}
        predictions = raw_report_baseline(self.cases, "pilot_v1")
        with self.assertRaisesRegex(ValueError, "Case dataset version"):
            evaluate(self.cases, references, predictions, [])

    def test_old_experiment_configuration_rejected_before_calls(self):
        with self.assertRaisesRegex(ValueError, "Case dataset version"):
            build_plan(load_json(ROOT / "evals/experiments.json"), self.cases, [])

    def test_model_input_does_not_include_form_metadata_or_possible_annotations(self):
        case = {**self.cases[0], "reference_summary": "Never show this clinician answer to the model."}
        payload = case_input(case)
        self.assertNotIn("dataset_version", payload)
        self.assertNotIn("patient_details", payload)
        self.assertNotIn("reference_summary", payload)
        self.assertIn("80 kg", payload["current_transcript"])
        recovering = next(case for case in self.cases if case["case_id"] == "SYN006")
        self.assertNotIn("68 kg", recovering["prior_information"])
        self.assertIn("68 kg", recovering["current_transcript"])

    def test_patient_workbook_import_ignores_archived_responses_and_rejects_formulas(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("Optional openpyxl not installed")
        book = MagicMock()
        book.sheetnames = ["Patient review", "Original v2"]
        sheet = MagicMock()
        sheet.max_row = len(self.rows)
        cells = [tuple(SimpleNamespace(value=value, data_type="s", row=i) for value in row)
                 for i, row in enumerate(self.rows, 1)]
        sheet.iter_rows.side_effect = lambda **kwargs: iter(cells)
        book.__getitem__.side_effect = lambda name: sheet if name == "Patient review" else self.fail("Read archived answers")
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "mock-source.bin"
            source.write_bytes(b"Mocked reader source for hash, not an authored workbook")
            with patch("openpyxl.load_workbook", return_value=book):
                reference = import_workbook(source, self.cases, [])
                self.assertEqual(reference["dataset_version"], "patient_review_v2.1")
                self.assertEqual(len(reference["cases"]), 10)
                self.assertEqual(reference["queues"], [])
                self.assertTrue(all(row["approved"] is False for row in reference["cases"]))
                with self.assertRaisesRegex(ValueError, "Missing required review tabs"):
                    import_workbook(source, self.cases, [{"queue_id": "Q01"}])
                cells[4][1].data_type = "f"
                with self.assertRaisesRegex(ValueError, "Formula in source/answer area"):
                    import_workbook(source, self.cases, [])
        self.assertTrue(book.close.called)

    def test_attention_rows_and_complete_answers_are_approved_deterministically(self):
        from evals.annotations import approve_complete, parse_attention_rows
        groups = load_json(ROOT / "docs/clinician-annotation/patient_review_v2_1_queues.json")
        source = {c["case_id"]: c for c in self.cases}
        rows = []
        for number, group in enumerate(groups, 1):
            rows += [[f"Group {number} / {group['queue_id']}", "Separate scenario."],
                     ["Is ranking meaningful here?", "Yes"], ["Why / missing context?", "Bed state unknown."]]
            for rank, case_id in enumerate(group["case_ids"], 1):
                case = source[case_id]
                rows += [[case_id, f"ETA {case['eta_minutes']} minutes at latest report"],
                         ["Age / reported sex / weight", case["patient_details"]],
                         ["T=0\nEMS Field Report", case["current_transcript"]],
                         ["Attention order", "Cannot assess" if rank == 3 else str(rank)],
                         ["Reason for position / tie / limits", f"Reason {rank}."]]
        queues = parse_attention_rows(rows, groups, self.cases, "Dr Test")
        self.assertEqual(queues[0]["order"], {"SYN001": 1, "SYN004": 2, "SYN007": "Unsure"})
        self.assertEqual(queues[0]["status"], "Done")
        self.assertIn("SYN004: Reason 2.", queues[0]["explanation"])
        changed = [list(row) for row in rows]
        changed[5][1] = changed[5][1].replace("46", "47")
        with self.assertRaisesRegex(ValueError, "Attention source"):
            parse_attention_rows(changed, groups, self.cases)

        self.rows.insert(0, ["Reviewer alias", "Dr Test"])
        for label, value in (("EMS to ED Handoff", "Handoff."), ("Meaningful change?", "No earlier report"),
                             ("Evidence / missing information", "Gaps."), ("Preparation needed?", "Routine"),
                             ("Review status", "Done")):
            self.set_answer(label, value)
        reference = {"dataset_version": "patient_review_v2.1", "cases": parse_patient_rows(self.rows, self.cases),
                     "queues": queues}
        skipped = approve_complete(reference)
        self.assertTrue(reference["cases"][0]["approved"])
        self.assertEqual(reference["cases"][0]["reviewer_id"], "Dr Test")
        self.assertEqual(reference["cases"][0]["rubric_version"], "patient_review_v2.1")
        self.assertFalse(reference["cases"][1]["approved"])
        self.assertTrue(all(q["approved"] for q in queues))
        self.assertEqual(len(skipped), 9)


if __name__ == "__main__":
    unittest.main()
