"""Read the revised ESI workbook without modifying it or converting v2 labels.

Only the current Patient review sheet supplies cases/labels. Archived answers,
ranking exercises and form feedback never become text features or ESI targets.
"""

import hashlib
import re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from .contracts import CHANGE, number, text

DATASET_VERSION = "patient_review_v2.1"
IMPORT_VERSION = "esi_annotation_import_v1"
SOURCE_FIELDS = ("case_id", "age_sex_weight", "prior_information", "current_transcript",
                 "elapsed_minutes", "eta_minutes")
ANSWER_LABELS = {"EMS to ED Handoff": "reference_summary", "Meaningful change?": "meaningful_change",
                 "Direction (if Yes)": "change_direction", "Evidence / missing information": "evidence_missing_information",
                 "ESI level (1–5)": "esi_level", "ESI reason / limits": "esi_reason",
                 "Review status": "review_status"}


def sha256_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_patient_sheet(path):
    """Bound actual rows, not unreliable/missing worksheet dimension metadata."""
    from openpyxl import load_workbook
    with ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 20_000_000:
            raise ValueError("Annotation workbook exceeds the 20 MB uncompressed limit")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        matches = [s for s in workbook if s.title.casefold() == "patient review"]
        if len(matches) != 1:
            raise ValueError("Expected one Patient review sheet")
        rows = []
        # This form has two core columns. Other columns are audit notes, never features.
        for i, row in enumerate(matches[0].iter_rows(), 1):
            if i > 2000 or len(row) > 64:
                raise ValueError("Patient review sheet exceeds bounded import size")
            if any(c.data_type in ("f", "e") for c in row):
                raise ValueError(f"Formula/error cell in Patient review row {i}; formulas are never evaluated")
            if any(isinstance(c.value, str) and len(c.value) > 50_000 for c in row):
                raise ValueError("Annotation cell exceeds text length limit")
            rows.append([c.value for c in row])
        return rows, workbook.sheetnames
    finally:
        workbook.close()


def parse_patient_rows(rows):
    versions = [row[1] for row in rows if len(row) > 1 and row[0] == "Dataset version"]
    if versions != [DATASET_VERSION]:
        raise ValueError("Expected exactly one patient_review_v2.1 dataset marker; do not import v2 attention labels as ESI")
    sources, references, locations, extras = [], [], {}, []
    current_source = current_reference = None
    seen_fields = set()
    for row_number, row in enumerate(rows, 1):
        if not row:
            continue
        label = text(row[0])
        value = row[1] if len(row) > 1 else None
        for column, extra in enumerate(row[2:], 3):
            if text(extra):
                extras.append({"row": row_number, "column": column, "value": extra})
        match = re.fullmatch(r"Case \d+ / (SYN\d+)", label)
        if match:
            cid = match[1]
            if cid in locations:
                raise ValueError(f"Duplicate case ID: {cid}")
            eta = re.fullmatch(r"ETA (\d+) minutes at latest report", text(value))
            if not eta:
                raise ValueError(f"Missing/invalid ETA header for {cid}")
            current_source = {"case_id": cid, "age_sex_weight": None, "prior_information": "",
                              "current_transcript": None, "elapsed_minutes": 0, "eta_minutes": int(eta[1])}
            current_reference = {"case_id": cid, **dict.fromkeys(ANSWER_LABELS.values()),
                                 "approved": False, "label_status": "provisional_single_reviewer"}
            sources.append(current_source)
            references.append(current_reference)
            locations[cid] = {"header": f"B{row_number}"}
            seen_fields = set()
            continue
        if current_source is None:
            continue
        if label == "Age / reported sex / weight":
            field = "age_sex_weight"
        elif label in ("T=0\nEMS Field Report", "T=0\nUpdated EMS Field Report"):
            field = "current_transcript"
        elif re.fullmatch(r"T-\d+ min\nEMS Field Report", label):
            field = "prior_information"
            current_source["elapsed_minutes"] = int(re.search(r"T-(\d+)", label)[1])
        elif label in ANSWER_LABELS:
            field = ANSWER_LABELS[label]
        elif label:
            raise ValueError(f"Unrecognized v2.1 case field at A{row_number}; version layout changes explicitly")
        else:
            if text(value):
                extras.append({"row": row_number, "column": 2, "value": value})
            continue
        if field in seen_fields:
            raise ValueError(f"Duplicate {field} for {current_source['case_id']}")
        seen_fields.add(field)
        locations[current_source["case_id"]][field] = f"B{row_number}"
        if field in SOURCE_FIELDS:
            current_source[field] = value
        else:
            # Preserve raw labels; do not map Improvement/Deterioration into v2 enums.
            current_reference[field] = value
    if not sources:
        raise ValueError("No v2.1 patient cases found")
    issues = []
    for source, reference in zip(sources, references):
        cid = source["case_id"]
        for field in ("age_sex_weight", "current_transcript"):
            if not isinstance(source[field], str) or not source[field].strip():
                issues.append({"case_id": cid, "field": field, "problem": "missing/invalid source text"})
        if not isinstance(source["prior_information"], str):
            issues.append({"case_id": cid, "field": "prior_information", "problem": "invalid source text"})
        for field in ANSWER_LABELS.values():
            if field != "change_direction" and not text(reference[field]):
                issues.append({"case_id": cid, "field": field, "problem": "missing annotation"})
        value = reference["esi_level"]
        if not number(value, integer=True, minimum=1) or value > 5:
            issues.append({"case_id": cid, "field": "esi_level", "problem": "ESI must be numeric 1–5; no inferred target"})
        else:
            reference["esi_level"] = int(value)
        if reference["review_status"] != "Done":
            issues.append({"case_id": cid, "field": "review_status", "problem": "annotation not Done"})
        if reference["meaningful_change"] not in CHANGE:
            issues.append({"case_id": cid, "field": "meaningful_change", "problem": "unknown label"})
        if not source["prior_information"] and reference["meaningful_change"] != "No earlier report":
            issues.append({"case_id": cid, "field": "meaningful_change", "problem": "no earlier source report to compare"})
    return {"kind": IMPORT_VERSION, "dataset_version": DATASET_VERSION, "cases": sources,
            "references": references, "source_cells": locations, "extra_notes": extras, "issues": issues,
            "warnings": ["Done is completion, not adjudication. Labels remain provisional; no clinical approval inferred.",
                         "This previously inspected pilot is development data, not a fresh independent clinical holdout.",
                         "Doctor summaries/explanations and archived sheets are excluded from feature inputs.",
                         "Meaningful-change labels and directions are preserved but are not this experiment's ESI target."],
            "esi_counts": dict(Counter(str(r["esi_level"]) for r in references)),
            "absent_esi_levels": [i for i in range(1, 6) if not any(r["esi_level"] == i for r in references)]}


def import_workbook(path):
    before = sha256_file(path)
    rows, sheets = read_patient_sheet(path)
    result = parse_patient_rows(rows)
    if sha256_file(path) != before:
        raise ValueError("Workbook changed during import; retry from a fixed snapshot")
    result.update({"source_sha256": before, "source_sheet": "Patient review",
                   "ignored_sheets": [s for s in sheets if s.casefold() != "patient review"]})
    return result
